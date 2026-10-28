"""Builds excel/RINS-Hub.xlsx, the starting workbook for the hub's Excel link.

Run:  python3 excel/build_workbook.py
The hub address and public key are read from config.js (both are public).
"""
import pathlib
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = pathlib.Path(__file__).resolve().parent.parent
cfg = (ROOT / "config.js").read_text()
HUB_URL = re.search(r'SUPABASE_URL:\s*"([^"]+)"', cfg).group(1)
HUB_KEY = re.search(r'SUPABASE_KEY:\s*"([^"]+)"', cfg).group(1)

NAVY, LIME, MUTED, TINT = "002A6D", "A3E400", "5B6888", "EEF9CF"
H1 = Font(name="Calibri", size=18, bold=True, color=NAVY)
H2 = Font(name="Calibri", size=12, bold=True, color=NAVY)
LABEL = Font(name="Calibri", size=11, bold=True)
NOTE = Font(name="Calibri", size=10, color=MUTED)
WRAP = Alignment(wrap_text=True, vertical="top")
KEYFILL = PatternFill("solid", fgColor=TINT)

wb = Workbook()

# ---------------- Start ----------------
ws = wb.active
ws.title = "Start"
ws.sheet_properties.tabColor = LIME
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 60
ws.column_dimensions["C"].width = 3
ws.column_dimensions["D"].width = 26
ws["A1"] = "RINS Hub · Excel link"
ws["A1"].font = H1
ws["A2"] = "A copy of the hub's data for analysis and reports, plus sheets for sending new entries to the hub."
ws["A2"].font = NOTE

rows = [
    (4, "Hub address", HUB_URL, "HubUrl"),
    (5, "Public key", HUB_KEY, "HubPublicKey"),
    (6, "Your Excel key", "Not set on this computer. Click Set my key.", "KeyStatus"),
    (8, "Last refresh", "", "LastRefresh"),
    (9, "Key name", "", "KeyName"),
    (10, "Access", "", "KeyAccess"),
    (11, "Last send", "", "LastSend"),
]
for r, label, value, name in rows:
    ws.cell(r, 1, label).font = LABEL
    c = ws.cell(r, 2, value)
    c.number_format = "@"
    if name in ("HubUrl", "HubPublicKey"):
        c.font = NOTE
    if name == "KeyStatus":
        c.fill = KEYFILL
    wb.defined_names[name] = DefinedName(name, attr_text=f"Start!$B${r}")

steps = [
    ("How to use it", H2),
    ("Refresh from hub: brings in the latest products, branches, stock, sales and (if your key allows) transactions, accounts and profit and loss. Each refresh replaces the data sheets.", None),
    ("Send to hub: saves the rows you typed on New sales, New stock moves and New money. Each row gets a Status: ✓ saved, or ✗ with the reason to fix. Saved rows are never sent twice.", None),
    ("Edit or delete past entries on the hub website, so there is one clear history.", None),
    ("Build your own pivots and charts on new sheets. Point them at the tables (tblSales, tblStock, tblTransactions, tblProfitandloss, …) and they update on every refresh. Don't type in the data sheets; a refresh overwrites them.", None),
    ("Your Excel key is saved on this computer, not in this file. You can share the file safely; each person sets their own key.", None),
    ("First time on this computer", H2),
    ("1. File > Save As, choose Excel Macro-Enabled Workbook (*.xlsm) and save it in your OneDrive.", None),
    ("2. Press Alt+F11 to open the code window, then File > Import File… and choose RinsHub.bas. Close the code window and save.", None),
    ("3. Press Alt+F8, choose SetupButtons and click Run. The buttons appear.", None),
    ("4. Click Set my key and paste the key from the hub (Company > Excel link). Then click Refresh from hub.", None),
]
r = 13
for text, font in steps:
    c = ws.cell(r, 1, text)
    if font:
        c.font = font
        r += 1
        continue
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    c.alignment = WRAP
    ws.row_dimensions[r].height = 32
    r += 1


# ---------------- entry sheets ----------------
def entry_sheet(title, table, cols, widths, hint, validations):
    s = wb.create_sheet(title)
    s.sheet_properties.tabColor = NAVY
    s["A1"] = title
    s["A1"].font = H1
    s.row_dimensions[2].height = 32
    s["A3"] = hint
    s["A3"].font = NOTE
    for i, (col, w) in enumerate(zip(cols, widths), start=1):
        s.cell(4, i, col)
        s.column_dimensions[s.cell(4, i).column_letter].width = w
    last = s.cell(4, len(cols)).column_letter
    t = Table(displayName=table, ref=f"A4:{last}5")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    s.add_table(t)
    s.freeze_panes = "A5"
    # Date column, text columns, hidden reference column
    for rr in range(5, 1001):
        s.cell(rr, 1).number_format = "yyyy-mm-dd"
    ref_letter = last
    s.column_dimensions[ref_letter].hidden = True
    for col_name, formula in validations:
        idx = cols.index(col_name) + 1
        letter = s.cell(4, idx).column_letter
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=False)
        dv.add(f"{letter}5:{letter}1000")
        s.add_data_validation(dv)
    dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                        showErrorMessage=True, errorTitle="Date", error="Type a date, for example 2026-10-28.")
    dv.add("A5:A1000")
    s.add_data_validation(dv)
    return s


BR = 'INDIRECT("tblBranches[Name]")'
PR = 'INDIRECT("tblProducts[Name]")'
entry_sheet(
    "New sales", "tblNewSales",
    ["Date", "Branch", "Product", "Qty", "Note", "Status", "Ref"],
    [13, 28, 28, 8, 30, 46, 10],
    "One row per product sold at a branch. The price is the product's shelf price on the hub. Branch and Product lists fill in after the first refresh.",
    [("Branch", BR), ("Product", PR)],
)
entry_sheet(
    "New stock moves", "tblNewMoves",
    ["Date", "Type", "Product", "From", "To", "Qty", "Note", "Status", "Ref"],
    [13, 14, 28, 24, 24, 8, 26, 46, 10],
    "Stock in: To only. Transfer and Return: From and To. Adjustment: From only, and the Qty can be negative.",
    [("Type", '"Stock in,Transfer,Return,Adjustment"'), ("Product", PR), ("From", BR), ("To", BR)],
)
entry_sheet(
    "New money", "tblNewMoney",
    ["Date", "Type", "Account", "To account", "Category", "Amount", "Paid to / from", "Branch",
     "Description", "Reference", "Receipt link", "Status", "Ref"],
    [13, 13, 14, 14, 30, 11, 22, 22, 26, 14, 26, 46, 10],
    "Money out / Money in need a Category. Transfer needs To account and no Category. Amounts in BHD. Needs a key with finance access.",
    [("Type", '"Money out,Money in,Transfer"'), ("Account", 'INDIRECT("tblAccounts[Account]")'),
     ("To account", 'INDIRECT("tblAccounts[Account]")'), ("Category", 'INDIRECT("tblCategories[Category]")'),
     ("Branch", BR)],
)

# ---------------- data sheets (filled by Refresh) ----------------
for name in ["Stock", "Sales", "Stock moves", "Products", "Branches", "Transactions", "Accounts",
             "Profit and loss", "Categories"]:
    s = wb.create_sheet(name)
    s["A1"] = "Click Refresh from hub on the Start sheet to fill this sheet."
    s["A1"].font = NOTE

out = ROOT / "excel" / "RINS-Hub.xlsx"
wb.save(out)
print("saved", out)
